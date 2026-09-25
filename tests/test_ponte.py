"""Le regole del ponte fra la libreria a cartelle e il portale (libreria_ponte.py).

La prima regola dell'utente: le funzioni che c'erano non si tolgono e non si
cambiano. Le altre discendono da li'.
"""
import os

import library
import library_db
import libreria_ponte
import store

from conftest import UNO, DUE, login, state, upload


def _library_id(client):
    code = client.get('/api/library/key').get_json()['code']
    return library_db.get_library_by_code(code)['id']


def _codice(client):
    return client.get('/api/library/key').get_json()['code']


def _script(app, code):
    """Come import_books.py: nessun login, solo il cookie con il codice."""
    client = app.test_client()
    client.set_cookie(library.COOKIE_NAME, code.replace('-', ''))
    return client


# ---------------------------------------------------------------------------
# Accanto, non al posto
# ---------------------------------------------------------------------------

def test_una_sola_libreria_su_libreria(app, client):
    pagina = client.get('/libreria')
    assert pagina.status_code == 200
    html = pagina.get_data(as_text=True)
    assert '/static/library.js' in html and 'I miei libri' in html        # la libreria unica
    # il vecchio indirizzo delle cartelle porta alla libreria unica
    vecchio = client.get('/libreria/cartelle')
    assert vecchio.status_code in (301, 302) and vecchio.headers['Location'].endswith('/libreria')
    # la pagina di prima resta come rete di sicurezza, identica
    classica = client.get('/libreria/classica')
    assert classica.status_code == 200
    assert 'btn-importa' in classica.get_data(as_text=True)


def test_senza_login_niente_cartelle(app):
    anonimo = app.test_client()
    for pagina in ('/libreria', '/libreria/cartelle', '/libreria/classica'):
        risposta = anonimo.get(pagina)
        assert risposta.status_code in (301, 302) and '/login' in risposta.headers['Location'], pagina
    assert anonimo.get('/api/library').status_code == 401
    assert anonimo.post('/api/library/folders', json={'name': 'x'}).status_code == 401


def test_una_libreria_per_account(app, client, other_client, pdf_file):
    upload(client, pdf_file)
    assert len(state(client)['books']) == 1
    assert state(other_client)['books'] == []
    assert _library_id(client) == _library_id(login(app, UNO))    # stesso account, stessa libreria


# ---------------------------------------------------------------------------
# Ogni libro delle cartelle e' un libro della libreria di sempre
# ---------------------------------------------------------------------------

def test_un_libro_caricato_si_legge_nel_portale(app, client, pdf_file):
    risposta = upload(client, pdf_file)
    assert risposta.status_code == 201
    libro = risposta.get_json()['book']
    assert libro['job_id'] and libro['read_url'] == f"/leggi/{libro['job_id']}"

    job = store.get_job(libro['job_id'], owner=UNO)
    assert job['status'] == 'completed' and job['provider'] == 'importato'
    record = library_db.get_book_record(_library_id(client), libro['id'])
    assert os.path.realpath(job['output_path']) == os.path.realpath(library_db.abs_path(record['file_path']))

    # compare nella libreria di sempre e il lettore riceve il file
    ids = [j['task_id'] for j in client.get('/api/jobs?limit=50').get_json()['jobs']]
    assert libro['job_id'] in ids
    file = client.get(f"/api/book/{libro['job_id']}/file")
    assert file.status_code == 200 and file.data[:5] == b'%PDF-'
    assert client.get(libro['read_url']).status_code == 200


def test_nessun_audiolibro_accodato(app, client, pdf_file, epub_file):
    upload(client, pdf_file)
    upload(client, epub_file)
    assert store.db().execute('SELECT COUNT(*) FROM audiolibri').fetchone()[0] == 0


def test_i_libri_che_ci_sono_gia_compaiono_senza_copie(app, client, tmp_path, pdf_file):
    percorso = str(tmp_path / 'gia_in_libreria.pdf')
    with open(pdf_file, 'rb') as a, open(percorso, 'wb') as b:
        b.write(a.read())
    store.create_job(id='11111111-2222-3333-4444-555555555555', owner=UNO,
                     original_filename='Gia in libreria.pdf', input_path=None, output_path=percorso,
                     output_filename='Gia in libreria.pdf', file_type='pdf',
                     source_lang='Inglese', target_lang='Italiano', provider='anthropic', model='claude',
                     status='completed', progress=1.0, status_text='Completato!')
    library_id = _library_id(client)
    assert libreria_ponte.sincronizza(library_id) == (1, 0)
    assert libreria_ponte.sincronizza(library_id) == (0, 0)          # mai due volte

    libri = state(client)['books']
    assert len(libri) == 1 and libri[0]['job_id'] == '11111111-2222-3333-4444-555555555555'
    record = library_db.get_book_record(library_id, libri[0]['id'])
    assert library_db.abs_path(record['file_path']) == percorso        # stesso file, nessuna copia
    assert len(os.listdir(tmp_path)) >= 1 and os.path.exists(percorso)


def test_caricamento_e_sincronizzazione_non_fanno_doppioni(app, client, pdf_file, epub_file):
    upload(client, pdf_file)
    upload(client, epub_file)
    assert libreria_ponte.sincronizza(_library_id(client)) == (0, 0)
    assert len(state(client)['books']) == 2
    assert len(client.get('/api/jobs?limit=50').get_json()['jobs']) == 2


# ---------------------------------------------------------------------------
# Eliminare
# ---------------------------------------------------------------------------

def test_il_cestino_non_toglie_nulla_dalla_libreria_di_sempre(app, client, pdf_file):
    libro = upload(client, pdf_file).get_json()['book']
    client.post('/api/library/trash', json={'book_ids': [libro['id']]})
    job = store.get_job(libro['job_id'], owner=UNO)
    assert job is not None and os.path.exists(job['output_path'])


def test_eliminare_per_sempre_dal_cestino_toglie_anche_dalla_libreria(app, client, pdf_file):
    libro = upload(client, pdf_file).get_json()['book']
    percorso = store.get_job(libro['job_id'], owner=UNO)['output_path']
    client.post('/api/library/trash', json={'book_ids': [libro['id']]})
    client.post('/api/library/trash/purge', json={'book_ids': [libro['id']]})
    assert store.get_job(libro['job_id'], owner=UNO) is None
    assert not os.path.exists(percorso)


def test_eliminato_da_libreria_sparisce_anche_dalle_cartelle(app, client, pdf_file):
    libro = upload(client, pdf_file).get_json()['book']
    assert client.delete(f"/api/libro/{libro['job_id']}").get_json() == {'ok': True}   # funzione di sempre
    assert libreria_ponte.sincronizza(_library_id(client)) == (0, 1)
    assert state(client)['books'] == []
    assert client.get('/api/library/trash').get_json()['items'] == []


def test_una_traduzione_in_corso_non_si_elimina(app, client, pdf_file):
    libro = upload(client, pdf_file).get_json()['book']
    store.update_job(libro['job_id'], status='running')
    client.post('/api/library/trash', json={'book_ids': [libro['id']]})
    client.post('/api/library/trash/purge', json={'book_ids': [libro['id']]})
    assert store.get_job(libro['job_id'], owner=UNO)['status'] == 'running'


# ---------------------------------------------------------------------------
# Il codice libreria: solo per import_books.py
# ---------------------------------------------------------------------------

def test_il_codice_apre_solo_le_chiamate_dello_script(app, client, pdf_file):
    script = _script(app, _codice(client))
    assert script.post('/api/library/key', json={'code': _codice(client)}).status_code == 200
    albero = script.post('/api/library/folders/tree', json={'paths': ['Marketing']})
    assert albero.status_code == 200
    cartella = albero.get_json()['folders']['Marketing']
    assert upload(script, pdf_file, folder_id=cartella).status_code == 201
    libro = state(client)['books'][0]
    assert libro['folder_id'] == cartella and libro['job_id']

    # tutto il resto vuole il login
    assert script.get('/api/library').status_code == 401
    assert script.get(f"/api/library/books/{libro['id']}/download").status_code == 401
    assert script.post('/api/library/trash', json={'book_ids': [libro['id']]}).status_code == 401
    assert script.post('/api/library/key/rotate').status_code == 401
    assert script.get('/libreria/cartelle').status_code in (301, 302)


def test_un_codice_sbagliato_non_apre_nulla(app, client):
    sbagliato = _script(app, 'ABCDEFGHJKMNPQRSTVWX')
    risposta = sbagliato.post('/api/library/key', json={'code': 'ABCDEFGHJKMNPQRSTVWX'})
    assert risposta.status_code == 404
    assert sbagliato.post('/api/library/folders/tree', json={'paths': ['x']}).status_code == 401


def test_il_codice_di_un_account_non_piu_autorizzato_non_vale(app, client, monkeypatch):
    code = _codice(client)
    monkeypatch.setenv('ALLOWED_EMAILS', DUE)
    assert _script(app, code).post('/api/library/key', json={'code': code}).status_code == 404


def test_con_il_login_non_si_apre_la_libreria_di_un_altro(app, client, other_client):
    altrui = _codice(other_client)
    assert client.post('/api/library/key', json={'code': altrui}).status_code == 404
    assert _library_id(client) != _library_id(other_client)


def test_il_codice_si_cambia_solo_dopo_il_login(app, client):
    vecchio = _codice(client)
    nuovo = client.post('/api/library/key/rotate').get_json()['code']
    assert nuovo != vecchio
    assert _script(app, vecchio).post('/api/library/key', json={'code': vecchio}).status_code == 404
    assert _script(app, nuovo).post('/api/library/key', json={'code': nuovo}).status_code == 200


# ---------------------------------------------------------------------------
# Titoli: quelli lasciati da scanner ed editor non valgono il nome del file
# ---------------------------------------------------------------------------

def test_i_titoli_spazzatura_lasciano_il_posto_al_nome_del_file():
    casi = {
        ('CamScanner 07-02-2020 18.02.04', 'Copywriting Secrets & Tactics.pdf'): 'Copywriting Secrets & Tactics',
        ('Preview of “WIZARD of ads cropped.pdf”', 'The Wizard of Ads.pdf'): 'The Wizard of Ads',
        ('How to write Letters_234x156.pdf, page 1-320', 'How to Write Sales Letters That Sell.pdf'):
            'How to Write Sales Letters That Sell',
        ('高績效心智', 'How to Do Better Creative Work.epub'): 'How to Do Better Creative Work',
        ('Microsoft Word - bozza3.docx', 'Il mio libro.pdf'): 'Il mio libro',
        ('', 'Senza titolo.epub'): 'Senza titolo',
    }
    for (titolo, nome), atteso in casi.items():
        assert library.clean_title(titolo, nome) == atteso, (titolo, nome)
    # i titoli veri restano, anche in altri alfabeti se il file non e' in latino
    assert library.clean_title('Influence: The Psychology of Persuasion', 'influence.epub') == \
        'Influence: The Psychology of Persuasion'
    assert library.clean_title('Documenting Software Architectures', 'dsa.pdf') == 'Documenting Software Architectures'
    assert library.clean_title('高績效心智', '高績效心智.epub') == '高績效心智'


# ---------------------------------------------------------------------------
# Libreria unica: tutti i formati di prima, e i libri tradotti che compaiono subito
# ---------------------------------------------------------------------------

def test_docx_txt_e_md_si_aggiungono_come_con_importa(app, client, tmp_path):
    import docx
    documento = docx.Document()
    documento.add_paragraph('Capitolo uno. Il copywriting spiegato semplice.')
    percorso_docx = str(tmp_path / 'Il mio libro (z-library.sk, 1lib.sk).docx')
    documento.save(percorso_docx)
    percorso_txt = tmp_path / 'appunti_tradotto_Italiano.txt'
    percorso_txt.write_text('Un testo semplice.', encoding='utf-8')
    percorso_md = tmp_path / 'scaletta.md'
    percorso_md.write_text('# Scaletta', encoding='utf-8')

    titoli = {}
    for percorso in (percorso_docx, str(percorso_txt), str(percorso_md)):
        risposta = upload(client, percorso)
        assert risposta.status_code == 201, risposta.get_json()
        libro = risposta.get_json()['book']
        assert libro['job_id'] and libro['read_url'] is None        # si scarica, non si apre nel lettore
        job = store.get_job(libro['job_id'], owner=UNO)
        assert job['provider'] == 'importato' and job['file_type'] == libro['file_type']
        titoli[libro['file_type']] = libro['title']
    assert titoli == {'docx': 'Il mio libro', 'txt': 'appunti', 'md': 'scaletta'}


def test_i_libri_docx_del_portale_entrano_nella_libreria(app, client, tmp_path):
    percorso = str(tmp_path / 'tradotto.docx')
    import docx
    docx.Document().save(percorso)
    store.create_job(id='22222222-3333-4444-5555-666666666666', owner=UNO,
                     original_filename='Manuale_tradotto_Italiano.docx', input_path=None, output_path=percorso,
                     output_filename='Manuale_tradotto_Italiano.docx', file_type='docx',
                     source_lang='Inglese', target_lang='Italiano', provider='anthropic', model='claude',
                     status='completed', progress=1.0, status_text='Completato!')
    assert libreria_ponte.sincronizza(_library_id(client)) == (1, 0)
    libro = state(client)['books'][0]
    assert libro['title'] == 'Manuale' and libro['file_type'] == 'docx' and libro['read_url'] is None


def test_la_pagina_sincronizza_subito_quando_una_traduzione_finisce(app, client, tmp_path, pdf_file):
    percorso = str(tmp_path / 'appena_tradotto.pdf')
    with open(pdf_file, 'rb') as a, open(percorso, 'wb') as b:
        b.write(a.read())
    store.create_job(id='33333333-4444-5555-6666-777777777777', owner=UNO,
                     original_filename='Appena tradotto.pdf', input_path=None, output_path=percorso,
                     output_filename='Appena tradotto.pdf', file_type='pdf',
                     source_lang='Inglese', target_lang='Italiano', provider='anthropic', model='claude',
                     status='completed', progress=1.0, status_text='Completato!')
    risposta = client.post('/api/library/sync')
    assert risposta.status_code == 200 and risposta.get_json()['indicizzati'] == 1
    assert [b['job_id'] for b in state(client)['books']] == ['33333333-4444-5555-6666-777777777777']
    # col solo codice libreria (lo script) la sincronizzazione non si chiama
    assert _script(app, _codice(client)).post('/api/library/sync').status_code == 401
